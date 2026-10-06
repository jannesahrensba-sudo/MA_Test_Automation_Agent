"""
Builds the Excel import template for existing test cases (docs/vorlagen/testfall-import-vorlage.xlsx).

The code lists (process teams, processes, ways, objects, process steps, process profiles) come from the mock data of
the app (webapp/localService/mainService/data), so the dropdowns match the mockup. The single example test case is
fictional and marked as such. Nothing else is invented: unknown values stay empty and are imported as open assignments.

Usage (from zstc-testautomation/):  pip install openpyxl && python3 tools/templates/import_template.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "zstc-testautomation/webapp/localService/mainService/data"
TARGET = ROOT / "docs/vorlagen/testfall-import-vorlage.xlsx"
ROWS = 500  # rows with dropdowns

HEADER_FILL = PatternFill("solid", fgColor="0A6ED1")
HEADER_FONT = Font(bold=True, color="FFFFFF")
REQUIRED_FILL = PatternFill("solid", fgColor="0854A0")
EXAMPLE_FONT = Font(italic=True, color="6A6D70")
WRAP = Alignment(wrap_text=True, vertical="top")

OBJECT_DE = {
    "SERVICE_REQUEST": "Service Request",
    "SERVICE_QUOTATION": "Angebot (Service Quotation)",
    "SERVICE_ORDER": "Serviceauftrag (Service Order)",
    "SERVICE_CONFIRMATION": "Rückmeldung (Service Confirmation)",
    "BILLING_DOC_REQUEST": "Fakturaanforderung (Billing Document Request)",
    "BILLING_DOCUMENT": "Faktura (Billing Document)",
    "ACCOUNTING_DOCUMENT": "Buchhaltungsbeleg (FI)",
    "SERVICE_CONTRACT": "Servicevertrag (wird ermittelt)",
}
START_OBJECTS = ["SERVICE_REQUEST", "SERVICE_QUOTATION", "SERVICE_ORDER", "SERVICE_CONFIRMATION", "BILLING_DOC_REQUEST", "BILLING_DOCUMENT"]
END_OBJECTS = ["SERVICE_REQUEST", "SERVICE_QUOTATION", "SERVICE_ORDER", "SERVICE_CONFIRMATION", "BILLING_DOC_REQUEST", "BILLING_DOCUMENT", "ACCOUNTING_DOCUMENT"]


def load(name):
    return json.loads((DATA / f"{name}.json").read_text(encoding="utf-8"))


# (column, description, required, list name or None, width)
CASE_COLUMNS = [
    ("ExterneTestfallID", "ID im bisherigen Werkzeug bzw. in Ihrer Excel-Datei", True, None, 18),
    ("Titel", "Kurztitel (max. 80 Zeichen)", True, None, 40),
    ("Prozessteam", "Team, das den Testfall verantwortet; leer = offen", False, "Prozessteams", 18),
    ("Prozess", "Geschäftsprozess; leer = offen", False, "Prozesse", 12),
    ("Weg", "Weg durch den Prozess (Prozessvariante); leer = offen", False, "Wege", 18),
    ("StartAb", "Start des Laufs; leer = dort, wo das Team in den Weg einsteigt", False, "Startobjekte", 22),
    ("LaufBis", "Ende des Laufs (Endobjekt); leer = Faktura", False, "Endobjekte", 22),
    ("VorgaengerTestfall", "Testfall, dessen Belege übernommen werden (bei Start ab Rückmeldung, Fakturaanforderung oder Faktura)", False, None, 20),
    ("Prozessprofil", "Profil der Pflichtfelder und Vorbelegungen; leer = aus dem Gerätetyp", False, "Prozessprofile", 16),
    ("Voraussetzungen", "Vorbedingungen und Übergaben in Worten", False, None, 36),
    ("Testdaten", "Feld=Wert, getrennt durch | (Feldnamen siehe Blatt Codes)", False, None, 46),
    ("ErwarteterNettowert", "erwarteter Nettowert in EUR, nur wenn bekannt", False, None, 14),
    ("FachlichVerantwortlich", "Benutzer mit Rolle Process Owner", False, None, 18),
    ("Quelle", "Datei, Blatt, Zeile oder Version der Vorlage", False, None, 24),
    ("OffenePunkte", "Was unklar ist (wird als offene Zuordnung übernommen, nicht geraten)", False, None, 36),
]
STEP_COLUMNS = [
    ("ExterneTestfallID", "Verweis auf das Blatt Testfälle", True, None, 18),
    ("SchrittNr", "laufende Nummer", True, None, 10),
    ("ProzessschrittID", "Prozessschritt; leer = manueller Schritt ohne Prozessbezug", False, "Prozessschritte", 16),
    ("Aktion", "Was wird getan?", True, None, 44),
    ("ErwartetesErgebnis", "Was muss herauskommen?", True, None, 44),
    ("Beleg", "Business Object, das der Schritt erzeugt oder prüft", False, "Endobjekte", 22),
    ("Pruefwerte", "Feld=Wert, getrennt durch | (z. B. NetValue=114.00)", False, None, 30),
    ("Hinweis", "Freitext", False, None, 30),
]

TEST_DATA_FIELDS = [
    ("SoldToParty", "Kunde (Auftraggeber, z. B. Hausverwaltung)"),
    ("ServiceRefFunctionalLocation", "Nutzeinheit (Technischer Platz)"),
    ("ServiceReferenceEquipment", "Gerät (Equipment)"),
    ("ServiceRequestReporter", "Meldender (Ansprechpartner des Kunden)"),
    ("ServiceRequestDescription", "Problembeschreibung (max. 40 Zeichen)"),
    ("ServiceDocumentPriority", "Priorität 1, 3, 5 oder 9"),
    ("RespyMgmtServiceTeam", "Serviceteam"),
    ("ServiceContract", "Servicevertrag (Weg 3)"),
    ("ServiceProduct", "Leistung"),
    ("ServiceDuration", "Einsatzdauer in Stunden"),
    ("ServicePart", "Ersatzteil"),
    ("ServicePartQuantity", "Menge Ersatzteil"),
    ("NetAmountTolerance", "Toleranz Nettowert in EUR"),
]


def header(ws, columns):
    for index, (name, description, required, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.fill = REQUIRED_FILL if required else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        note = ws.cell(row=2, column=index, value=("Pflicht: " if required else "") + description)
        note.font = Font(italic=True, size=9, color="32363A")
        note.alignment = WRAP
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "A3"


def add_list(ws, title, rows, start_column):
    """writes a code list (code, text) and returns the absolute range of the codes"""
    ws.cell(row=1, column=start_column, value=title).font = Font(bold=True)
    ws.cell(row=2, column=start_column, value="Code").font = Font(bold=True)
    ws.cell(row=2, column=start_column + 1, value="Bedeutung").font = Font(bold=True)
    for offset, (code, text) in enumerate(rows):
        ws.cell(row=3 + offset, column=start_column, value=code)
        ws.cell(row=3 + offset, column=start_column + 1, value=text)
    ws.column_dimensions[get_column_letter(start_column)].width = 24
    ws.column_dimensions[get_column_letter(start_column + 1)].width = 46
    letter = get_column_letter(start_column)
    return f"Codes!${letter}$3:${letter}${2 + len(rows)}"


def validations(ws, columns, ranges):
    for index, (name, _, _, list_name, _) in enumerate(columns, start=1):
        if not list_name:
            continue
        validation = DataValidation(type="list", formula1=ranges[list_name], allow_blank=True)
        # warning instead of stop: values of the old tool may be entered and are then imported as open assignment
        validation.errorStyle = "warning"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        validation.errorTitle = "Wert nicht in der Liste"
        validation.error = "Der Wert ist nicht in der Codeliste (Blatt Codes). Er wird beim Import nicht übernommen, die Zuordnung bleibt offen."
        validation.promptTitle = name
        validation.prompt = "Auswahl aus Blatt Codes; leer lassen, wenn unbekannt."
        letter = get_column_letter(index)
        validation.add(f"{letter}3:{letter}{ROWS}")
        ws.add_data_validation(validation)


def main():
    teams = load("ProcessTeamVH")
    processes = load("BusinessProcessVH")
    variants = load("ProcessVariantVH")
    profiles = load("ProcessProfileVH")
    steps = sorted((s for s in load("ProcessStep") if s.get("IsActiveEntity")), key=lambda s: s["Sequence"])

    wb = Workbook()

    # --- instructions
    guide = wb.active
    guide.title = "Anleitung"
    lines = [
        ("Importvorlage für bestehende Testfälle – Service-to-Cash Test Automation Assistant (Mockup)", Font(bold=True, size=14)),
        ("Zielsystem: SAP S/4HANA 2025 (Releasezyklus mit Feature Pack Stacks). Die Codes in den Auswahllisten stammen aus dem Mockup; Kunden-, Geräte- und Belegdaten sind fiktiv.", None),
        ("", None),
        ("So füllen Sie die Vorlage aus", Font(bold=True)),
        ("1. Blatt „Testfälle“: eine Zeile je Testfall (Kopf). Pflichtspalten sind dunkelblau.", None),
        ("2. Blatt „Testschritte“: eine Zeile je Testschritt, verknüpft über die ExterneTestfallID.", None),
        ("3. Auswahllisten nutzen (Blatt „Codes“). Ist ein Wert unbekannt, lassen Sie das Feld leer und beschreiben Sie es unter OffenePunkte.", None),
        ("4. Die Zeile BEISPIEL-001 ist fiktiv und dient nur der Orientierung – vor dem Import löschen.", None),
        ("Alternativ können Sie Ihre vorhandene Excel-Datei unverändert bereitstellen; die Spalten werden dann nach dieser Vorlage zugeordnet.", None),
        ("", None),
        ("Regeln bei der Übernahme (nichts wird erfunden)", Font(bold=True)),
        ("• Fehlt Team, Prozess oder Weg oder ist der Wert unbekannt, bleibt das Feld leer: Die Zuordnung wird „offen“ mit Hinweis, was fehlt.", None),
        ("• Ist das Team nur aus dem Kontext ableitbar, wird es als „Annahme“ gekennzeichnet und später fachlich bestätigt.", None),
        ("• Testschritte ohne Prozessschritt bleiben manuelle Schritte und zählen nicht zur Schrittabdeckung.", None),
        ("• Freigaben aus dem bisherigen Werkzeug werden nicht übernommen: Jeder Testfall wird als Version 1 neu validiert und vom Process Owner freigegeben.", None),
        ("• Startet ein Testfall mitten im Prozess (z. B. „ab Fakturaanforderung“ im Prozessteam New End to End Prozess), braucht er einen Vorgänger-Testfall, dessen Belege er übernimmt.", None),
        ("• Die Teststufe (Teilprozess oder End-to-End) wird aus Start, Ende und Teamwechseln abgeleitet, nicht übernommen.", None),
        ("• Die Zuordnung zu Releases erfolgt über den Release-Umfang (Prozessteam × Prozess), nicht je Testfall.", None),
    ]
    for row, (text, font) in enumerate(lines, start=1):
        cell = guide.cell(row=row, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font
    guide.column_dimensions["A"].width = 140

    # --- code lists
    codes = wb.create_sheet("Codes")
    ranges = {
        "Prozessteams": add_list(codes, "Prozessteams", [(t["ProcessTeam"], f"{t['ProcessTeamName']} ({t['ProcessArea']})") for t in teams], 1),
        "Prozesse": add_list(codes, "Prozesse", [(p["ProcessID"], p["ProcessName"] + ("" if p["PilotScope"] == "PILOT" else " (später)")) for p in processes], 4),
        "Wege": add_list(codes, "Wege (Reparaturprozess)", [(v["Variant"], v["VariantName"]) for v in variants], 7),
        "Startobjekte": add_list(codes, "Start ab", [(code, OBJECT_DE[code]) for code in START_OBJECTS], 10),
        "Endobjekte": add_list(codes, "Lauf bis / Beleg", [(code, OBJECT_DE[code]) for code in END_OBJECTS], 13),
        "Prozessschritte": add_list(
            codes,
            "Prozessschritte (Reparaturprozess)",
            [(s["StepID"], f"{s['StepName']} · {s.get('ResponsibleTeam') or 'Team offen'}" + ("" if s.get("PilotScope") == "PILOT" else " (später)")) for s in steps],
            16,
        ),
        "Prozessprofile": add_list(codes, "Prozessprofile", [(p["ProcessProfile"], p["ProcessProfileName"]) for p in profiles], 19),
    }
    add_list(codes, "Testdaten: Feldnamen", TEST_DATA_FIELDS, 22)

    # --- test cases
    cases = wb.create_sheet("Testfälle", 1)
    header(cases, CASE_COLUMNS)
    example = {
        "ExterneTestfallID": "BEISPIEL-001",
        "Titel": "BEISPIEL (fiktiv, vor dem Import löschen): HKV-Störung ohne Angebot",
        "Prozessteam": "PT-REPARATUR",
        "Prozess": "SRV-REP",
        "Weg": "W1_REQUEST",
        "StartAb": "SERVICE_REQUEST",
        "LaufBis": "BILLING_DOCUMENT",
        "Prozessprofil": "MD_HKV_STOER",
        "Voraussetzungen": "Gerät ist in der Nutzeinheit eingebaut",
        "Testdaten": "ServiceReferenceEquipment=HKV-0815-031|ServiceRequestReporter=MD-CP-1001",
        "Quelle": "Mockup-Beispiel",
    }
    for index, (name, *_rest) in enumerate(CASE_COLUMNS, start=1):
        cell = cases.cell(row=3, column=index, value=example.get(name))
        cell.font = EXAMPLE_FONT
    validations(cases, CASE_COLUMNS, ranges)
    cases.auto_filter.ref = f"A1:{get_column_letter(len(CASE_COLUMNS))}1"

    # --- test steps
    step_sheet = wb.create_sheet("Testschritte", 2)
    header(step_sheet, STEP_COLUMNS)
    example_steps = [
        ["BEISPIEL-001", 1, "REP-010", "Service Request mit Kunde, Gerät und Problembeschreibung anlegen", "Service Request angelegt", "SERVICE_REQUEST", "", ""],
        ["BEISPIEL-001", 2, "", "Telefonische Rückfrage beim Bewohner", "Termin vereinbart", "", "", "manueller Schritt ohne Prozessschritt"],
        ["BEISPIEL-001", 3, "REP-100", "Faktura erzeugen", "Faktura mit erwartetem Nettowert", "BILLING_DOCUMENT", "NetValue=108.00", ""],
    ]
    for row, values in enumerate(example_steps, start=3):
        for column, value in enumerate(values, start=1):
            cell = step_sheet.cell(row=row, column=column, value=value if value != "" else None)
            cell.font = EXAMPLE_FONT
    validations(step_sheet, STEP_COLUMNS, ranges)
    step_sheet.auto_filter.ref = f"A1:{get_column_letter(len(STEP_COLUMNS))}1"

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TARGET)
    print(f"written {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
